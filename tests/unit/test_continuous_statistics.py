import csv
import json
from typing import get_args, get_origin

import pytest
from openpyxl import load_workbook
from pydantic import BaseModel, ValidationError

from deet.custom_hierarchical import (
    _build_dynamic_models_from_schema,
    _build_dynamic_pipeline_for_study_type,
    _load_prompt_schema,
    _serialize_value,
    _write_dynamic_study_outputs,
    write_hierarchical_prompts_csv,
)
from deet.hierarchical_mvp.cochrane_postprocessing import postprocess_cochrane_outcomes
from deet.hierarchical_mvp.continuous_statistics import (
    ReportedStatistic,
    expand_statistics_columns,
    revman_statistics,
)
from deet.hierarchical_mvp.rct_registry import RCT_MODEL_MODULES
from deet.hierarchical_mvp.utils import (
    _write_tables,
    export_animal_csv,
    export_cochrane_csv,
    export_csv,
    export_obesity_csv,
)


def test_named_statistic_serialization_preserves_measure_and_context():
    statistic = ReportedStatistic(
        measure="Confidence interval",
        value="8-17",
        basis="Change from baseline",
        details="95% CI for mean change",
    )
    assert _serialize_value([statistic]) == [statistic.model_dump()]
    assert _serialize_value(statistic)["measure"] == "Confidence interval"


@pytest.mark.parametrize("study_type", RCT_MODEL_MODULES)
def test_all_continuous_schemas_require_named_statistic_lists(study_type):
    cls = RCT_MODEL_MODULES[study_type].Continuous_Outcome
    fields = {
        name: field
        for name, field in cls.model_fields.items()
        if name.endswith("_statistics")
    }
    assert len(fields) in (2, 4)
    for field in fields.values():
        assert get_args(field.annotation) == (ReportedStatistic,)
        assert field.is_required()
    assert not {
        "group_mean",
        "group_standard_deviation",
        "mean",
        "standard_deviation",
        "group_a_mean",
        "group_b_mean",
        "mean_a",
        "mean_b",
    }.intersection(cls.model_fields)
    with pytest.raises(ValidationError):
        cls.model_validate({"group_mean": "12", "group_standard_deviation": "2"})


@pytest.mark.parametrize("study_type", RCT_MODEL_MODULES)
def test_prompt_csv_statistic_options_remain_editable(tmp_path, study_type):
    path = write_hierarchical_prompts_csv(study_type, tmp_path / "prompts.csv")
    schema = _load_prompt_schema(path)
    measure = next(
        field
        for field in schema["ReportedStatistic"]
        if field["attribute"] == "measure"
    )
    assert measure["datatype"] == "str"
    measure["prompt"] = "Extract trimmed mean or median absolute deviation."
    models = _build_dynamic_models_from_schema(schema)
    continuous = models["Continuous_Outcome"]
    for name, field in continuous.model_fields.items():
        if not name.endswith("_statistics"):
            continue
        statistic_cls = get_args(field.annotation)[0]
        assert statistic_cls.model_fields["measure"].description == measure["prompt"]
        instance = statistic_cls(measure="Trimmed mean", value="12")
        assert instance.measure == "Trimmed mean"
        assert "enum" not in statistic_cls.model_json_schema()["properties"]["measure"]
        assert field.default_factory is list


@pytest.mark.parametrize("study_type", RCT_MODEL_MODULES)
def test_old_mean_sd_prompt_csvs_are_explicitly_rejected(tmp_path, study_type):
    path = write_hierarchical_prompts_csv(study_type, tmp_path / "prompts.csv")
    schema = _load_prompt_schema(path)
    schema["Continuous_Outcome"] = [
        field
        for field in schema["Continuous_Outcome"]
        if not field["attribute"].endswith("_statistics")
    ]
    schema["Continuous_Outcome"].append(
        {"attribute": "group_mean", "datatype": "str", "prompt": "Mean"}
    )
    with pytest.raises(ValueError, match="Regenerate the prompt CSV"):
        _build_dynamic_pipeline_for_study_type(study_type, schema)


@pytest.mark.parametrize(
    ("measure", "basis", "details", "column", "expected"),
    [
        ("Mean", "Endpoint", "", "Mean", "12"),
        ("Median", "Endpoint", "", "Mean", ""),
        ("Mean change", "Change from baseline", "", "Mean", ""),
        ("Mean", "Change from baseline", "", "Mean", ""),
        ("Mean", "NR", "", "Mean", ""),
        ("Mean", "Endpoint", "Adjusted", "Mean", ""),
        ("Standard deviation", "Endpoint", "", "SD", "12"),
        ("Standard error", "Endpoint", "", "SE", "12"),
        ("Interquartile range", "Endpoint", "", "SD", ""),
        ("Variance", "Endpoint", "", "Variance", "12"),
    ],
)
def test_revman_projection_never_relabels_statistics(
    measure,
    basis,
    details,
    column,
    expected,
):
    statistic = ReportedStatistic(
        measure=measure,
        value="12",
        basis=basis,
        details=details,
    )
    row = revman_statistics([statistic], [statistic])
    assert row[column] == expected
    assert json.loads(row["group_summary_statistics"]) == [statistic.model_dump()]
    assert json.loads(row["group_variability_statistics"]) == [statistic.model_dump()]


def test_revman_does_not_choose_arbitrarily_between_multiple_means():
    statistics = [
        ReportedStatistic(measure="Mean", value="10", basis="Endpoint"),
        ReportedStatistic(measure="Mean", value="12", basis="Endpoint"),
    ]
    row = revman_statistics(statistics, [])
    assert row["Mean"] == ""
    assert '"10"' in row["group_summary_statistics"]
    assert '"12"' in row["group_summary_statistics"]


@pytest.mark.parametrize("study_type", ["CochraneRCT", "CochraneRCT_pairwise"])
def test_cochrane_exports_and_postprocessing_preserve_non_mean_statistics(
    tmp_path,
    study_type,
):
    median = ReportedStatistic(measure="Median", value="12", basis="Endpoint")
    iqr = ReportedStatistic(
        measure="Interquartile range",
        value="8-17",
        basis="Endpoint",
        details="Q1-Q3 for median",
    )
    if study_type.endswith("_pairwise"):
        payload = {
            "outcome_name": "MADRS",
            "arm_a": "Treatment",
            "arm_b": "Control",
            "sample_size_a": "40",
            "sample_size_b": "39",
            "group_a_summary_statistics": [median.model_dump()],
            "group_a_variability_statistics": [iqr.model_dump()],
            "group_b_summary_statistics": [
                ReportedStatistic(
                    measure="Mean", value="16", basis="Endpoint"
                ).model_dump(),
            ],
            "group_b_variability_statistics": [],
        }
    else:
        payload = {
            "outcome_name": "MADRS",
            "group_label": "Treatment",
            "group_N": "40",
            "group_summary_statistics": [median.model_dump()],
            "group_variability_statistics": [iqr.model_dump()],
        }
    model = RCT_MODEL_MODULES[study_type].Continuous_Outcome.model_validate(payload)
    rows = model.to_csv_rows(study="study")
    assert rows[0]["Mean"] == ""
    assert rows[0]["SD"] == ""
    assert json.loads(rows[0]["group_summary_statistics"]) == [median.model_dump()]
    assert json.loads(rows[0]["group_variability_statistics"]) == [iqr.model_dump()]
    if study_type.endswith("_pairwise"):
        assert rows[1]["Mean"] == "16"
        assert rows[1]["group_variability_statistics"] == "[]"
    path = tmp_path / "study.json"
    path.write_text(
        json.dumps(
            {
                "study_characteristics": {"study": "study"},
                "continuous_outcomes": [model.model_dump()],
            }
        ),
        encoding="utf-8",
    )
    output = postprocess_cochrane_outcomes(path)
    with output.open(encoding="utf-8-sig", newline="") as source:
        processed = list(csv.DictReader(source))
    assert len(processed) == len(rows)
    _, expected_rows = expand_statistics_columns(
        list(rows[0]), [dict(row) for row in rows],
        single=not study_type.endswith("_pairwise"),
    )
    for expected, actual in zip(expected_rows, processed, strict=True):
        assert all(actual[key] == value for key, value in expected.items())


@pytest.mark.parametrize("prefix", ["group", "group_a", "group_b"])
def test_statistics_columns_preserve_all_values_and_avoid_collisions(prefix):
    field = f"{prefix}_summary_statistics"
    collision = f"{field}_1_value"
    statistics = [
        ReportedStatistic(
            measure="Mean change", value="-8.6", basis="Change from baseline"
        ),
        ReportedStatistic(measure="Median", value="12", details="Adjusted"),
    ]
    rows = [
        {field: [], collision: "custom", f"{field}_1_VALUE_2": "also custom"},
        {field: json.dumps([item.model_dump() for item in statistics])},
        {field: statistics},
    ]
    headers, expanded = expand_statistics_columns(list(rows[0]), rows)
    assert len(headers) == len({name.casefold() for name in headers})
    assert field not in headers
    assert expanded[0][collision] == "custom"
    assert expanded[0][f"{field}_1_VALUE_2"] == "also custom"
    assert expanded[0][f"{field}_1_value_3"] == ""
    for row in expanded[1:]:
        assert row[f"{field}_1_measure"] == "Mean change"
        assert row[f"{field}_1_value_3"] == "-8.6"
        assert row[f"{field}_1_basis"] == "Change from baseline"
        assert row[f"{field}_1_details"] == ""
        assert row[f"{field}_2_measure"] == "Median"
        assert row[f"{field}_2_details"] == "Adjusted"
    assert rows[0][field] == []
    assert isinstance(rows[1][field], str)


def test_empty_statistic_tables_have_flat_headers():
    headers, rows = expand_statistics_columns(["group_summary_statistics"], [])
    assert headers == [
        "group_summary_statistics_1_measure",
        "group_summary_statistics_1_value",
        "group_summary_statistics_1_basis",
        "group_summary_statistics_1_details",
    ]
    assert rows == []


@pytest.mark.parametrize("study_type", RCT_MODEL_MODULES)
@pytest.mark.parametrize("dynamic", [False, True])
@pytest.mark.parametrize("kind", ["summary", "variability"])
def test_statistic_cardinality_is_enforced(tmp_path, study_type, dynamic, kind):
    module = RCT_MODEL_MODULES[study_type]
    payload = _sample_statistics_model(module.Continuous_Outcome).model_dump()
    cls = module.Continuous_Outcome
    if dynamic:
        path = write_hierarchical_prompts_csv(study_type, tmp_path / "prompts.csv")
        models = _build_dynamic_models_from_schema(_load_prompt_schema(path))
        cls = models["Continuous_Outcome"]
    field = next(
        name for name in cls.model_fields if name.endswith(f"_{kind}_statistics")
    )
    statistic = {"measure": "Other", "value": "12"}
    payload[field] = [statistic, statistic]
    if study_type.endswith("_pairwise"):
        assert len(getattr(cls.model_validate(payload), field)) == 2
    else:
        assert cls.model_json_schema()["properties"][field]["maxItems"] == 1
        with pytest.raises(ValidationError, match="at most 1 item"):
            cls.model_validate(payload)
        study_payload = _sample_statistics_model(module.Study).model_dump()
        study_payload["continuous_outcomes"] = [payload]
        study_cls = models["Study"] if dynamic else module.Study
        with pytest.raises(ValidationError, match="at most 1 item"):
            study_cls.model_validate(study_payload)


def test_single_statistic_columns_have_exactly_eight_fields():
    fields = ["group_summary_statistics", "group_variability_statistics"]
    statistic = {"measure": "Mean", "value": "12", "basis": "Endpoint", "details": ""}
    original = [
        {"group_label": "A", fields[0]: [statistic], fields[1]: []},
        {"group_label": "B", fields[0]: [], fields[1]: [statistic]},
    ]
    headers, rows = expand_statistics_columns(
        ["group_label", *fields], original, single=True
    )
    assert headers == ["group_label"] + [
        f"{field}_{key}"
        for field in fields
        for key in ("measure", "value", "basis", "details")
    ]
    assert len(rows) == 2
    assert rows[0]["group_summary_statistics_value"] == "12"
    assert rows[0]["group_variability_statistics_value"] == ""
    assert rows[1]["group_summary_statistics_value"] == ""
    assert rows[1]["group_variability_statistics_value"] == "12"
    empty_headers, empty_rows = expand_statistics_columns(
        ["group_label", *fields], [], single=True
    )
    assert empty_headers == headers
    assert empty_rows == []
    assert original[0][fields[0]] == [statistic]
    with pytest.raises(ValueError, match="separate outcome objects"):
        expand_statistics_columns(
            fields, [{fields[0]: [statistic, statistic]}], single=True
        )


def test_single_statistic_columns_avoid_custom_header_collisions():
    field = "group_summary_statistics"
    collision = f"{field}_VALUE"
    headers, rows = expand_statistics_columns(
        [field, collision],
        [{field: [{"measure": "Median", "value": "12"}], collision: "custom"}],
        single=True,
    )
    assert len(headers) == len({name.casefold() for name in headers})
    assert rows[0][collision] == "custom"
    assert rows[0][f"{field}_value_2"] == "12"


@pytest.mark.parametrize("study_type", ["CochraneRCT", "CochraneRCT_pairwise"])
def test_cochrane_dichotomous_only_postprocessing_preserves_export_mode(
    tmp_path, study_type,
):
    module = RCT_MODEL_MODULES[study_type]
    study = _sample_statistics_model(module.Study)
    study.dichotomous_outcomes.append(
        _sample_statistics_model(module.Dichotomous_Outcome)
    )
    path = tmp_path / "study.json"
    path.write_text(study.model_dump_json(), encoding="utf-8")
    output = postprocess_cochrane_outcomes(path)
    with output.open(encoding="utf-8-sig", newline="") as source:
        headers = csv.DictReader(source).fieldnames
    index = "_1" if study_type.endswith("_pairwise") else ""
    assert f"group_summary_statistics{index}_measure" in headers
    assert f"group_variability_statistics{index}_measure" in headers


def test_shared_writer_expands_statistics_without_touching_other_lists(tmp_path):
    statistic = {
        "measure": "Mean change",
        "value": "-8.6",
        "basis": "Change from baseline",
    }
    rows = [
        {"outcome": "empty", "group_summary_statistics": "[]", "notes": '["note"]'},
        {"outcome": "reported", "group_summary_statistics": [statistic, statistic]},
    ]
    _write_tables(
        {"outcomes": (list(rows[0]), rows)},
        tmp_path,
        "test",
        "",
        write_csv=True,
        write_xlsx=True,
    )
    with (tmp_path / "outcomes_test.csv").open(
        encoding="utf-8-sig", newline=""
    ) as source:
        reader = csv.DictReader(source)
        csv_rows = list(reader)
        headers = reader.fieldnames
    assert csv_rows[0]["notes"] == '["note"]'
    assert csv_rows[0]["group_summary_statistics_2_value"] == ""
    assert csv_rows[1]["group_summary_statistics_2_value"] == "-8.6"
    workbook = load_workbook(tmp_path / "study_test.xlsx", read_only=True)
    try:
        values = list(workbook["outcomes"].values)
        assert list(values[0]) == headers
        assert len(headers) == len(set(headers))
        assert (
            dict(zip(headers, values[2], strict=True))[
                "group_summary_statistics_2_value"
            ]
            == "-8.6"
        )
    finally:
        workbook.close()


def _sample_statistics_model(model_cls):
    payload = {}
    for name, field in model_cls.model_fields.items():
        annotation = field.annotation
        if get_origin(annotation) is list:
            payload[name] = []
        elif get_origin(annotation) is dict:
            payload[name] = {}
        elif isinstance(annotation, type) and issubclass(annotation, BaseModel):
            payload[name] = _sample_statistics_model(annotation)
        else:
            payload[name] = "NR"
    return model_cls.model_validate(payload)


@pytest.mark.parametrize("study_type", RCT_MODEL_MODULES)
@pytest.mark.parametrize("dynamic", [False, True])
def test_static_and_dynamic_statistic_exports(tmp_path, study_type, dynamic):
    module = RCT_MODEL_MODULES[study_type]
    study = _sample_statistics_model(module.Study)
    outcome = _sample_statistics_model(module.Continuous_Outcome)
    for name in type(outcome).model_fields:
        if name.endswith("_statistics"):
            setattr(
                outcome, name, [ReportedStatistic(measure="Trimmed mean", value="-8.6")]
            )
    study.continuous_outcomes.append(outcome)
    original = study.model_dump()
    if dynamic:
        path = write_hierarchical_prompts_csv(study_type, tmp_path / "prompts.csv")
        schema = _load_prompt_schema(path)
        measure = next(
            field
            for field in schema["ReportedStatistic"]
            if field["attribute"] == "measure"
        )
        measure["prompt"] = "Extract trimmed means without converting them."
        models = _build_dynamic_models_from_schema(schema)
        study = models["Study"].model_validate(original)
        _write_dynamic_study_outputs(
            study,
            schema,
            study_type,
            tmp_path,
            "example",
            "test",
            "",
            export_csv_files=True,
            export_xlsx_file=True,
            export_json_file=True,
        )
    else:
        exports = {
            "RCT": export_csv,
            "AnimalRCT": export_animal_csv,
            "CochraneRCT": export_cochrane_csv,
            "ObesityRCT": export_obesity_csv,
        }
        exports[study_type.removesuffix("_pairwise")](
            study, "example", tmp_path, "test", write_csv=True, write_xlsx=True
        )
    directory = tmp_path / "example"
    with (directory / "outcomes_test.csv").open(
        encoding="utf-8-sig", newline=""
    ) as source:
        reader = csv.DictReader(source)
        rows = list(reader)
        headers = reader.fieldnames
    assert len(headers) == len({name.casefold() for name in headers})
    assert not any(name.endswith("_statistics") for name in headers)
    assert any(
        row[key] == "-8.6"
        for row in rows
        for key in headers
        if key.endswith("_value")
    )
    if not study_type.endswith("_pairwise"):
        for kind in ("summary", "variability"):
            assert [
                key for key in headers
                if key.startswith(f"group_{kind}_statistics_")
            ] == [
                f"group_{kind}_statistics_{key}"
                for key in ("measure", "value", "basis", "details")
            ]
    workbook = load_workbook(directory / "study_test.xlsx", read_only=True)
    try:
        assert list(next(workbook["outcomes"].values)) == headers
    finally:
        workbook.close()
    assert study.model_dump() == original


def test_cochrane_postprocessing_keeps_legacy_pairwise_mean_sd_json(tmp_path):
    path = tmp_path / "legacy.json"
    path.write_text(
        json.dumps(
            {
                "continuous_outcomes": [
                    {
                        "outcome_name": "MADRS",
                        "arm_a": "Treatment",
                        "arm_b": "Control",
                        "mean_a": "10",
                        "mean_b": "20",
                        "sd_a": "2",
                        "sd_b": "3",
                        "ci_level_a": "0.95",
                        "ci_start_a": "8",
                        "ci_end_a": "12",
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    output = postprocess_cochrane_outcomes(path)
    with output.open(encoding="utf-8-sig", newline="") as source:
        rows = list(csv.DictReader(source))
    assert [row["Mean"] for row in rows] == ["10", "20"]
    assert [row["SD"] for row in rows] == ["2", "3"]
    assert rows[0]["CI level"] == "0.95"
    assert rows[0]["CI start"] == "8"
    assert rows[0]["CI end"] == "12"


@pytest.mark.parametrize("study_type", RCT_MODEL_MODULES)
def test_exported_prompt_csv_includes_statistic_list_types(tmp_path, study_type):
    path = write_hierarchical_prompts_csv(study_type, tmp_path / "prompts.csv")
    with path.open(encoding="utf-8-sig", newline="") as source:
        rows = list(csv.DictReader(source))
    statistics = [
        row
        for row in rows
        if row["class"] == "Continuous_Outcome"
        and row["attribute"].endswith("_statistics")
    ]
    assert statistics
    assert all(row["datatype"] == "list[ReportedStatistic]" for row in statistics)
    assert {
        row["attribute"] for row in rows if row["class"] == "ReportedStatistic"
    } == {
        "measure",
        "value",
        "basis",
        "details",
    }

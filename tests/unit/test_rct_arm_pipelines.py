import csv
import json
from types import SimpleNamespace
from typing import get_origin
from unittest.mock import Mock, patch

import pytest
from openpyxl import load_workbook
from pydantic import BaseModel

from deet.custom_hierarchical import (
    _build_dynamic_pipeline_for_study_type,
    _load_prompt_schema,
    _write_dynamic_study_outputs,
    write_hierarchical_prompts_csv,
)
from deet.hierarchical_mvp.cochrane_postprocessing import postprocess_cochrane_outcomes
from deet.hierarchical_mvp.outcome_prompts import (
    CONTINUOUS_STATISTICS_INSTRUCTIONS,
    SINGLE_ARM_OUTCOME_INSTRUCTIONS,
)
from deet.hierarchical_mvp.rct_registry import RCT_MODEL_MODULES
from deet.main_hierarchical import RCT_PIPELINES, extract, save_data
from export_to_llmextract import (
    STUDY_TYPE_MODULES,
    _module_classes,
    build_template_workbook,
)

ARM_NAMES = ["Drug A", "Drug B", "Placebo"]


@pytest.mark.parametrize("study_type", RCT_MODEL_MODULES)
def test_intervention_name_prompts_use_experimental_groups(tmp_path, study_type):
    path = write_hierarchical_prompts_csv(study_type, tmp_path / "prompts.csv")
    schema = _load_prompt_schema(path)
    name_fields = (
        [
            ("InductionIntervention", "intervention_to_induce_name"),
            ("AssessmentIntervention", "intervention_to_assess_name"),
        ]
        if study_type.startswith("Animal")
        else [
            (
                "Intervention",
                "arm" if study_type.startswith("Cochrane") else "group_name",
            )
        ]
    )
    for class_name, attribute in name_fields:
        field = next(
            field for field in schema[class_name] if field["attribute"] == attribute
        )
        assert field["prompt"].startswith(
            "Identify and extract the name of the experimental groups from this study."
        )


@pytest.mark.parametrize("study_type", RCT_MODEL_MODULES)
def test_outcome_prompts_are_group_neutral(tmp_path, study_type):
    path = write_hierarchical_prompts_csv(study_type, tmp_path / "prompts.csv")
    schema = _load_prompt_schema(path)
    for class_name in ("Dichotomous_Outcome", "Continuous_Outcome", "Other_Outcome"):
        for field in schema.get(class_name, []):
            prompt = field["prompt"].lower()
            assert "intervention group" not in prompt
            assert "intervention arm" not in prompt
    if not study_type.endswith("_pairwise"):
        for class_name in (
            "Dichotomous_Outcome",
            "Continuous_Outcome",
            "Other_Outcome",
        ):
            if class_name in schema:
                label = next(
                    field
                    for field in schema[class_name]
                    if field["attribute"] == "group_label"
                )
                assert "this group" in label["prompt"]
                assert "control" in label["prompt"]


def _sample(model_cls):
    payload = {}
    for name, field in model_cls.model_fields.items():
        annotation = field.annotation
        if get_origin(annotation) is list:
            payload[name] = []
        elif get_origin(annotation) is dict:
            payload[name] = {}
        elif isinstance(annotation, type) and issubclass(annotation, BaseModel):
            payload[name] = _sample(annotation)
        elif annotation is bool:
            payload[name] = False
        else:
            payload[name] = "NR"
    return model_cls.model_validate(payload)


def _study_payload(study_type):
    module = RCT_MODEL_MODULES[study_type]
    study = _sample(module.Study)
    payload = study.model_dump()
    _populate_arms(module, payload)
    _populate_outcomes(module, payload, pairwise=study_type.endswith("_pairwise"))
    return payload


def _populate_arms(module, payload):
    for list_name, class_name in (
        ("interventions", "Intervention"),
        ("assessment_interventions", "AssessmentIntervention"),
        ("induction_interventions", "InductionIntervention"),
    ):
        if list_name not in payload:
            continue
        arms = []
        for name in ARM_NAMES:
            arm = _sample(getattr(module, class_name)).model_dump()
            for key in (
                "group_name",
                "arm",
                "intervention_to_assess_name",
                "intervention_to_induce_name",
            ):
                if key in arm:
                    arm[key] = name
            arms.append(arm)
        payload[list_name] = arms


def _populate_outcomes(module, payload, *, pairwise):
    for list_name, class_name in (
        ("dichotomous_outcomes", "Dichotomous_Outcome"),
        ("continuous_outcomes", "Continuous_Outcome"),
        ("other_outcomes", "Other_Outcome"),
    ):
        if list_name not in payload:
            continue
        outcomes = []
        for index, name in enumerate(ARM_NAMES):
            outcome = _sample(getattr(module, class_name)).model_dump()
            outcome["outcome_name"] = (
                "MADRS" if class_name == "Continuous_Outcome" else class_name
            )
            if pairwise:
                if "group_labels" in outcome:
                    outcome["group_labels"] = {"group_a": name, "group_b": "Control"}
                if "arm_a" in outcome:
                    outcome.update(arm_a=name, arm_b="Control")
            else:
                outcome["group_label"] = name
            for field in outcome:
                if field.endswith("_summary_statistics"):
                    outcome[field] = [
                        {
                            "measure": "Mean",
                            "value": str(10 + index),
                            "basis": "Endpoint",
                            "details": "",
                        },
                        {
                            "measure": "Median",
                            "value": str(9 + index),
                            "basis": "Endpoint",
                            "details": "",
                        },
                        {
                            "measure": "Mean change",
                            "value": "-3",
                            "basis": "Change from baseline",
                            "details": "",
                        },
                    ]
                elif field.endswith("_variability_statistics"):
                    outcome[field] = [
                        {
                            "measure": "Standard deviation",
                            "value": "2",
                            "basis": "Endpoint",
                            "details": "",
                        },
                        {
                            "measure": "Interquartile range",
                            "value": "8-17",
                            "basis": "Endpoint",
                            "details": "Q1-Q3 for median",
                        },
                        {
                            "measure": "Confidence interval",
                            "value": "9-13",
                            "basis": "Endpoint",
                            "details": "95% CI for mean",
                        },
                        {
                            "measure": "Standard error",
                            "value": "0.4",
                            "basis": "Endpoint",
                            "details": "",
                        },
                    ]
            outcomes.append(outcome)
        payload[list_name] = outcomes


def _prediction(payload):
    return SimpleNamespace(
        **payload, flexible_outcomes=payload.get("other_outcomes", [])
    )


def _assert_xlsx_statistics(directory, rows):
    workbook = load_workbook(next(directory.glob("study*.xlsx")), read_only=True)
    try:
        values = list(workbook["outcomes"].values)
        assert len(values) == len(rows) + 1
        for csv_row, xlsx_row in zip(rows, values[1:], strict=True):
            for key, value in zip(values[0], xlsx_row, strict=True):
                if key.endswith("_statistics") and csv_row[key]:
                    assert json.loads(value) == json.loads(csv_row[key])
    finally:
        workbook.close()


def _assert_tables(directory, study_type):
    csv_path = next(directory.glob("outcomes*.csv"))
    with csv_path.open(encoding="utf-8-sig", newline="") as source:
        rows = list(csv.DictReader(source))
    cochrane = study_type.startswith("Cochrane")
    expected = 6 if cochrane else 9
    if study_type == "CochraneRCT_pairwise":
        expected *= 2
    assert len(rows) == expected
    for row in rows:
        for key, value in row.items():
            if key.endswith("_summary_statistics") and value:
                items = json.loads(value)
                assert [item["measure"] for item in items] == [
                    "Mean",
                    "Median",
                    "Mean change",
                ]
                assert items[2]["basis"] == "Change from baseline"
            elif key.endswith("_variability_statistics") and value:
                items = json.loads(value)
                assert [item["measure"] for item in items] == [
                    "Standard deviation",
                    "Interquartile range",
                    "Confidence interval",
                    "Standard error",
                ]
                assert items[2]["details"] == "95% CI for mean"
    if not study_type.endswith("_pairwise"):
        label = "Arm" if cochrane else "group_label"
        assert [row[label] for row in rows[:3]] == ARM_NAMES
        continuous = [
            row
            for row in rows
            if row.get("Outcome", row.get("outcome_name")) == "MADRS"
        ]
        statistics = [json.loads(row["group_summary_statistics"]) for row in continuous]
        assert [items[0]["value"] for items in statistics] == ["10", "11", "12"]
        assert all(items[1]["measure"] == "Median" for items in statistics)
        assert all(
            json.loads(row["group_variability_statistics"])[1]["value"] == "8-17"
            for row in continuous
        )
        if cochrane:
            assert [row["Mean"] for row in continuous] == ["10", "11", "12"]
        assert not any(key.startswith(("group_a_", "group_b_")) for key in rows[0])
    _assert_xlsx_statistics(directory, rows)


@pytest.mark.parametrize("study_type", RCT_MODEL_MODULES)
@pytest.mark.parametrize("flat_output", [False, True])
def test_standard_pipeline_and_exports(tmp_path, study_type, flat_output):
    payload = _study_payload(study_type)
    with patch("dspy.Predict", return_value=Mock(return_value=_prediction(payload))):
        study = extract(context="Synthetic three-arm study.", study_type=study_type)
    assert study.model_dump() == payload
    pipeline = RCT_PIPELINES[study_type]()
    assert (
        CONTINUOUS_STATISTICS_INSTRUCTIONS
        in pipeline.extract_continuous.signature.instructions
    )
    for name in ("extract_dichotomous", "extract_continuous", "extract_other"):
        if not hasattr(pipeline, name):
            continue
        instructions = getattr(pipeline, name).signature.instructions
        assert (SINGLE_ARM_OUTCOME_INSTRUCTIONS in instructions) == (
            not study_type.endswith("_pairwise")
        )
    save_data(
        study,
        ["three-arm.md"],
        str(tmp_path),
        study_type,
        export_csv_files=True,
        export_xlsx_file=True,
        export_json_file=True,
        flat_output=flat_output,
    )
    assert (
        json.loads(next(tmp_path.glob("*.json")).read_text(encoding="utf-8")) == payload
    )
    _assert_tables(tmp_path if flat_output else tmp_path / "three-arm", study_type)


@pytest.mark.parametrize("study_type", RCT_MODEL_MODULES)
@pytest.mark.parametrize("flat_output", [False, True])
def test_custom_pipeline_and_exports(tmp_path, study_type, flat_output):
    schema_path = write_hierarchical_prompts_csv(study_type, tmp_path / "prompts.csv")
    schema = _load_prompt_schema(schema_path)
    pipeline_cls = _build_dynamic_pipeline_for_study_type(study_type, schema)
    payload = _study_payload(study_type)
    with patch("dspy.Predict", return_value=Mock(return_value=_prediction(payload))):
        study = pipeline_cls()(context="Synthetic three-arm study.")
    for name in ("continuous_outcomes", "dichotomous_outcomes", "other_outcomes"):
        if name in payload:
            assert len(getattr(study, name)) == 3
            assert [item.model_dump() for item in getattr(study, name)] == payload[name]
    pipeline = pipeline_cls()
    assert (
        CONTINUOUS_STATISTICS_INSTRUCTIONS
        in pipeline.extract_continuous.signature.instructions
    )
    for name in ("extract_dichotomous", "extract_continuous", "extract_other"):
        predictor = getattr(pipeline, name)
        if predictor is not None:
            assert (
                SINGLE_ARM_OUTCOME_INSTRUCTIONS in predictor.signature.instructions
            ) == (not study_type.endswith("_pairwise"))
    output = tmp_path / "exports"
    output.mkdir()
    path = _write_dynamic_study_outputs(
        study,
        schema,
        study_type,
        output,
        "three-arm",
        "test",
        "",
        export_csv_files=True,
        export_xlsx_file=True,
        export_json_file=True,
        flat_output=flat_output,
    )
    exported = json.loads(path.read_text(encoding="utf-8"))["continuous_outcomes"]
    assert len(exported) == 3
    for actual, expected in zip(exported, payload["continuous_outcomes"], strict=True):
        for key in expected:
            if key.endswith("_statistics"):
                assert actual[key] == expected[key]
    table_dir = output if flat_output else output / "three-arm"
    with next(table_dir.glob("outcomes*.csv")).open(
        encoding="utf-8-sig", newline=""
    ) as source:
        rows = list(csv.DictReader(source))
    assert len(rows) == (6 if study_type.startswith("Cochrane") else 9)
    continuous = [row for row in rows if row["outcome_name"] == "MADRS"]
    for row, expected in zip(
        continuous,
        payload["continuous_outcomes"],
        strict=True,
    ):
        for key, value in expected.items():
            if key.endswith("_statistics"):
                assert json.loads(row[key]) == value
    if not study_type.endswith("_pairwise"):
        assert [row["group_label"] for row in rows[:3]] == ARM_NAMES
        continuous = [row for row in rows if row["outcome_name"] == "MADRS"]
        assert [
            json.loads(row["group_summary_statistics"])[0]["value"]
            for row in continuous
        ] == ["10", "11", "12"]
        assert all(
            json.loads(row["group_variability_statistics"])[2]["details"]
            == "95% CI for mean"
            for row in continuous
        )
    _assert_xlsx_statistics(table_dir, rows)


@pytest.mark.parametrize(
    "study_type", ["RCT", "CochraneRCT", "ObesityRCT", "AnimalRCT"]
)
def test_single_arm_rejects_old_pairwise_prompt_schema(tmp_path, study_type):
    path = write_hierarchical_prompts_csv(
        f"{study_type}_pairwise", tmp_path / "old.csv"
    )
    with pytest.raises(ValueError, match="Regenerate the prompt CSV"):
        _build_dynamic_pipeline_for_study_type(study_type, _load_prompt_schema(path))


@pytest.mark.parametrize("study_type", ["CochraneRCT", "CochraneRCT_pairwise"])
def test_cochrane_postprocessing_supports_both_shapes(tmp_path, study_type):
    path = tmp_path / "study.json"
    path.write_text(json.dumps(_study_payload(study_type)), encoding="utf-8")
    output = postprocess_cochrane_outcomes(path)
    with output.open(encoding="utf-8-sig", newline="") as source:
        rows = list(csv.DictReader(source))
    assert len(rows) == (12 if study_type.endswith("_pairwise") else 6)
    assert [row["Arm"] for row in rows[::2] if study_type.endswith("_pairwise")] == (
        ARM_NAMES * 2 if study_type.endswith("_pairwise") else []
    )


@pytest.mark.parametrize("study_type", RCT_MODEL_MODULES)
def test_template_workbooks_include_matching_schema(study_type):
    classes = _module_classes(STUDY_TYPE_MODULES[study_type])
    workbook = build_template_workbook(classes)
    try:
        rows = list(workbook["Continuous_Outcome"].values)
        fields = {row[0] for row in rows[1:]}
        assert ("group_label" in fields) == (not study_type.endswith("_pairwise"))
        assert "ReportedStatistic" in workbook.sheetnames
        assert any(row[1] == "list[ReportedStatistic]" for row in rows[1:])
    finally:
        workbook.close()

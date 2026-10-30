import csv
import json
from collections import defaultdict
from io import StringIO
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from deet import custom_hierarchical, main_hierarchical
from deet.hierarchical_mvp.programme_framework import (
    FRAMEWORK_COLUMNS,
    PROGRAMME_EXTRACTION_INSTRUCTIONS,
    QUALITATIVE_OUTCOME_INSTRUCTIONS,
    QUANTITATIVE_OUTCOME_INSTRUCTIONS,
    SDG_CATEGORIES,
    load_indicator_framework,
    resolve_framework_config,
    validate_classifications,
    validate_sdg_categories,
)
from deet.hierarchical_mvp.ProgrammeExtraction import ProgrammeExtractionPipeline
from deet.hierarchical_mvp.ProgrammeModel import (
    Intervention,
    Programme,
    Programme_Characteristics,
    Qualitative_Outcome,
    Quantitative_Outcome,
)
from deet.hierarchical_mvp.utils import export_programme_csv


def _programme() -> Programme:
    characteristics = Programme_Characteristics(
        countries="Kenya",
        responsible_author="UN Evaluation Office",
        budget="USD 2 million",
    )
    intervention = Intervention(
        name="Skills training",
        description="Training for local partners",
        time_frame="2024-2025",
        resources_used="Facilitators and training materials",
        intended_results=(
            "Aim: all partners improve delivery capacity. "
            "Actual: 80% improved; the aim was not fully met."
        ),
        source_grounding=(
            '"The programme provided skills training for local partners" '
            "(Chapter 2, Intervention Design, p. 8)."
        ),
    )
    outcome = Quantitative_Outcome(
        name="Improved delivery capacity",
        description="Partners reported stronger delivery capacity",
        intervention_name=intervention.name,
        unit_of_analysis="Partner organisations",
        SDG_categories="Partnerships for the goals",
        baseline_data="Before training, 20% had adequate capacity; partners lacked skills.",
        results="80% of partners improved delivery capacity",
        results_data_source=(
            'Quantitative: "80% of partners improved delivery capacity" '
            "(Chapter 3, Findings, p. 12)."
        ),
        barrier="Staff turnover",
        barrier_source=(
            '"Staff turnover limited the continuity of delivery" '
            "(Chapter 4, Challenges, p. 18)."
        ),
        enabler="Local ownership",
        enabler_source=(
            '"Local ownership helped sustain delivery" '
            "(Chapter 4, Findings, p. 19)."
        ),
        lessons_learned="Train multiple staff per partner",
        lessons_learned_source=(
            '"Training several staff in each partner organisation is important" '
            "(Chapter 5, Lessons, p. 24)."
        ),
        achievements="Delivery continued after handover",
        achievements_source=(
            '"Delivery continued after the handover" '
            "(Chapter 3, Findings, p. 15)."
        ),
        recommendation="Expand the training-of-trainers model",
        recommendation_source=(
            '"The training-of-trainers model should be expanded" '
            "(Chapter 6, Recommendations, p. 30)."
        ),
        baseline_data_source=(
            '"Before training, 20% had adequate capacity" '
            "(Chapter 3, Baseline, p. 11)."
        ),
    )
    return Programme(
        programme_characteristics=characteristics,
        interventions=[intervention],
        quantitative_outcomes=[outcome],
        qualitative_outcomes=[Qualitative_Outcome(
            **{
                **outcome.model_dump(exclude={"outcome_type", "baseline_data"}),
                "results": "Interviewed partners reported stronger delivery capacity",
                "results_data_source": (
                    '"Stronger delivery capacity" '
                    "(Chapter 3, Interviews, p. 13)."
                ),
            }
        )],
    )


@pytest.mark.parametrize("kind", ["quantitative", "qualitative"])
def test_outcome_results_preserved_in_serialization(kind):
    outcome = getattr(_programme(), f"{kind}_outcomes")[0]
    for row in (outcome.model_dump(), outcome.to_csv_row()):
        assert row["results"] == outcome.results
        assert row["outcome_type"] == kind
        assert "quantitative_outcome_results" not in row
        assert "qualitative_outcome_results" not in row
        assert ("baseline_data" in row) == (kind == "quantitative")
    with pytest.raises(ValidationError, match="outcome_type"):
        type(outcome).model_validate({
            **outcome.model_dump(), "outcome_type": "incorrect"
        })


@pytest.mark.parametrize("model", [Quantitative_Outcome, Qualitative_Outcome])
def test_custom_programme_prompts_include_outcome_results(model):
    rows = {
        row["attribute"]: row
        for row in custom_hierarchical.build_programme_hierarchical_prompt_rows()
        if row["class"] == model.__name__
    }

    for field in ("results", "outcome_type"):
        assert rows[field]["datatype"] == "str"
        assert rows[field]["prompt"] == model.model_fields[field].description
    for field in (
        "results_data_source", "barrier_source", "enabler_source",
        "lessons_learned_source", "achievements_source",
        "recommendation_source", "SDG_categories",
        "outcome_category_classification"
    ):
        assert rows[field]["datatype"] == "str"
        assert rows[field]["prompt"] == model.model_fields[field].description
    assert ("baseline_data" in rows) == (model is Quantitative_Outcome)
    if model is Quantitative_Outcome:
        assert list(rows).index("baseline_data") + 2 == list(rows).index("results")
        assert rows["baseline_data_source"]["prompt"] == (
            model.model_fields["baseline_data_source"].description
        )
    assert "intended_results" not in rows
    assert "sdg_belonging" not in rows
    assert "strategic_plan_belonging" not in rows
    assert "quantitative_outcome_results" not in rows
    assert "qualitative_outcome_results" not in rows
    intervention_rows = {
        row["attribute"]: row
        for row in custom_hierarchical.build_programme_hierarchical_prompt_rows()
        if row["class"] == "Intervention"
    }
    assert intervention_rows["intended_results"]["prompt"] == (
        Intervention.model_fields["intended_results"].description
    )
    assert intervention_rows["source_grounding"]["prompt"] == (
        Intervention.model_fields["source_grounding"].description
    )


@pytest.mark.parametrize("dynamic", [False, True])
@pytest.mark.parametrize("custom_framework", [False, True])
def test_programme_pipeline_passes_prior_outputs_to_later_steps(
    dynamic, custom_framework, tmp_path
):
    programme = _programme()
    if dynamic:
        schema = defaultdict(list)
        for row in custom_hierarchical.build_programme_hierarchical_prompt_rows():
            schema[row["class"]].append(
                {
                    "attribute": row["attribute"],
                    "prompt": row["prompt"],
                    "datatype": row["datatype"],
                }
            )
        dynamic_models = custom_hierarchical._build_dynamic_models_from_schema(schema)
        runtime_models = custom_hierarchical._ensure_programme_runtime_models(
            schema, dynamic_models
        )
        programme = runtime_models["Programme"].model_validate(programme.model_dump())
        pipeline_class = custom_hierarchical._build_dynamic_programme_pipeline(
            runtime_models
        )
    else:
        pipeline_class = ProgrammeExtractionPipeline
    framework_path = None
    if custom_framework:
        framework_path = tmp_path / "framework.csv"
        framework_path.write_text(
            ",".join(FRAMEWORK_COLUMNS)
            + "\nX.1,Local goal,Local indicator,Delivery capacity,Training alone\n",
            encoding="utf-8",
        )
        for outcome in (
            *programme.quantitative_outcomes, *programme.qualitative_outcomes
        ):
            outcome.outcome_category_classification = "X.1: Local indicator"
    pipeline = pipeline_class(classification_framework_path=framework_path)
    assert len(pipeline.predictors()) == 3
    assert pipeline.extract_programme_info.signature.instructions == (
        PROGRAMME_EXTRACTION_INSTRUCTIONS.strip()
    )
    assert pipeline.extract_quantitative.signature.instructions == (
        QUANTITATIVE_OUTCOME_INSTRUCTIONS.strip()
    )
    assert pipeline.extract_qualitative.signature.instructions == (
        QUALITATIVE_OUTCOME_INSTRUCTIONS.strip()
    )
    pipeline.extract_programme_info = lambda **kwargs: SimpleNamespace(
        programme_characteristics=programme.programme_characteristics,
        interventions=programme.interventions,
    )
    outcomes_calls = []
    qualitative_calls = []

    def extract_quantitative(**kwargs):
        outcomes_calls.append(kwargs)
        assert not qualitative_calls
        return SimpleNamespace(quantitative_outcomes=programme.quantitative_outcomes)

    def extract_qualitative(**kwargs):
        assert outcomes_calls
        qualitative_calls.append(kwargs)
        return SimpleNamespace(qualitative_outcomes=programme.qualitative_outcomes)

    pipeline.extract_quantitative = extract_quantitative
    pipeline.extract_qualitative = extract_qualitative

    result = pipeline(context="report")

    assert result.model_dump() == programme.model_dump()
    assert set(result.model_dump()) == {
        "programme_characteristics", "interventions",
        "quantitative_outcomes", "qualitative_outcomes"
    }
    assert outcomes_calls == [
        {
            "context": "report",
            "classification_framework": load_indicator_framework(framework_path),
            "programme_characteristics": programme.programme_characteristics,
            "interventions": programme.interventions,
        }
    ]
    assert qualitative_calls == [{
        **outcomes_calls[0], "quantitative_outcomes": programme.quantitative_outcomes
    }]
    pipeline.extract_quantitative = lambda **_kwargs: SimpleNamespace(
        quantitative_outcomes=programme.quantitative_outcomes
    )
    programme.quantitative_outcomes[0].SDG_categories = "SDG 17"
    with pytest.warns(UserWarning, match="SDG_categories"):
        cleaned = pipeline(context="report")
    assert cleaned.quantitative_outcomes[0].SDG_categories == ""
    programme.qualitative_outcomes[0].SDG_categories = "Unknown"
    programme.quantitative_outcomes[0].SDG_categories = "Partnerships for the goals"
    programme.quantitative_outcomes[0].outcome_category_classification = "unknown: Invented"
    with pytest.warns(UserWarning, match="removed invalid or duplicate") as warnings:
        cleaned = pipeline(context="report")
    assert len(warnings) == 2
    assert cleaned.quantitative_outcomes[0].outcome_category_classification == ""
    assert cleaned.qualitative_outcomes[0].SDG_categories == ""
    assert cleaned.quantitative_outcomes[0].results == (
        programme.quantitative_outcomes[0].results
    )
    valid_category = (
        "X.1: Local indicator" if custom_framework else
        "1.1: Percentage of pregnant women receiving at least four antenatal visits"
    )
    programme.quantitative_outcomes[0].outcome_category_classification = (
        f"unknown: Invented; {valid_category}; {valid_category}"
    )
    programme.quantitative_outcomes[0].SDG_categories = (
        "No poverty; SDG 17; Partnerships for the goals; No poverty"
    )
    with pytest.warns(UserWarning) as warnings:
        cleaned = pipeline(context="report")
    assert len(warnings) == 2
    row = cleaned.quantitative_outcomes[0].model_dump()
    assert row["outcome_category_classification"] == valid_category
    assert row["SDG_categories"] == "No poverty; Partnerships for the goals"


def test_main_extract_dispatches_programme_pipeline():
    expected = _programme()
    with patch.object(
        main_hierarchical, "ProgrammeExtractionPipeline"
    ) as pipeline_class:
        pipeline_class.return_value.return_value = expected

        result = main_hierarchical.extract("report", "Programme")

    pipeline_class.return_value.assert_called_once_with(context="report")
    assert result is expected


@pytest.mark.parametrize("dynamic", [False, True])
@pytest.mark.parametrize(
    ("quantitative", "qualitative"),
    [(True, True), (False, True), (True, False), (False, False)],
)
def test_pipeline_handles_empty_outcome_groups_and_cleans_before_qualitative(
    dynamic, quantitative, qualitative
):
    programme = _programme()
    pipeline_class = ProgrammeExtractionPipeline
    if dynamic:
        schema = defaultdict(list)
        for row in custom_hierarchical.build_programme_hierarchical_prompt_rows():
            schema[row["class"]].append(row)
        models = custom_hierarchical._ensure_programme_runtime_models(
            schema, custom_hierarchical._build_dynamic_models_from_schema(schema)
        )
        programme = models["Programme"].model_validate(programme.model_dump())
        pipeline_class = custom_hierarchical._build_dynamic_programme_pipeline(models)
    if not quantitative:
        programme.quantitative_outcomes = []
    if not qualitative:
        programme.qualitative_outcomes = []
    for outcome in programme.quantitative_outcomes:
        outcome.SDG_categories = "Unknown; No poverty"
    pipeline = pipeline_class()
    calls = []

    def extract_programme_info(**_kwargs):
        calls.append("programme")
        return SimpleNamespace(
            programme_characteristics=programme.programme_characteristics,
            interventions=programme.interventions,
        )

    def extract_quantitative(**_kwargs):
        calls.append("quantitative")
        return SimpleNamespace(quantitative_outcomes=programme.quantitative_outcomes)

    def extract_qualitative(**kwargs):
        calls.append("qualitative")
        assert kwargs["quantitative_outcomes"] == programme.quantitative_outcomes
        for outcome in kwargs["quantitative_outcomes"]:
            assert outcome.SDG_categories == "No poverty"
        return SimpleNamespace(qualitative_outcomes=programme.qualitative_outcomes)

    pipeline.extract_programme_info = extract_programme_info
    pipeline.extract_quantitative = extract_quantitative
    pipeline.extract_qualitative = extract_qualitative
    if quantitative:
        with pytest.warns(UserWarning, match="SDG_categories"):
            result = pipeline(context="report")
    else:
        result = pipeline(context="report")
    assert calls == ["programme", "quantitative", "qualitative"]
    assert result.model_dump() == programme.model_dump()


def test_dynamic_programme_requires_migration_of_old_outcome_schema():
    with pytest.raises(ValueError, match="Regenerate"):
        custom_hierarchical._ensure_programme_runtime_models(
            {"Outcome": []}, {}
        )


def test_dynamic_programme_without_container_rows_uses_custom_outcome_fields():
    schema = defaultdict(list)
    for row in custom_hierarchical.build_programme_hierarchical_prompt_rows():
        if row["class"] != "Programme":
            schema[row["class"]].append(row)
    schema["Qualitative_Outcome"].append({
        "attribute": "participant_group",
        "prompt": "Identify the group giving feedback.",
        "datatype": "str",
    })
    models = custom_hierarchical._ensure_programme_runtime_models(
        schema, custom_hierarchical._build_dynamic_models_from_schema(schema)
    )
    programme = models["Programme"].model_validate(_programme().model_dump())
    programme.qualitative_outcomes[0].participant_group = "Local partners"
    assert programme.model_dump()["qualitative_outcomes"][0]["participant_group"] == (
        "Local partners"
    )
    assert programme.quantitative_outcomes[0].outcome_type == "quantitative"
    assert programme.qualitative_outcomes[0].outcome_type == "qualitative"
    with pytest.raises(ValidationError, match="outcome_type"):
        models["Qualitative_Outcome"].model_validate({
            **programme.qualitative_outcomes[0].model_dump(),
            "outcome_type": "quantitative",
        })


def test_custom_programme_schema_builds_programme_pipeline():
    schema = defaultdict(list)
    for row in custom_hierarchical.build_programme_hierarchical_prompt_rows():
        schema[row["class"]].append(
            {
                "attribute": row["attribute"],
                "prompt": row["prompt"],
                "datatype": row["datatype"],
            }
        )

    pipeline_class = custom_hierarchical._build_dynamic_pipeline_for_study_type(
        "Programme", dict(schema)
    )

    assert pipeline_class.__name__ == "DynamicProgrammeExtractionPipeline"
    assert set(schema) == {
        "Programme_Characteristics", "Intervention", "Programme",
        "Quantitative_Outcome", "Qualitative_Outcome"
    }
    for class_name in ("Quantitative_Outcome", "Qualitative_Outcome"):
        outcome_fields = {row["attribute"] for row in schema[class_name]}
        assert {
            "barrier", "enabler", "lessons_learned", "achievements", "recommendation",
            "results", "outcome_type",
        } <= outcome_fields
        assert "outcome_name" not in outcome_fields


@pytest.mark.parametrize("dynamic", [False, True])
@pytest.mark.parametrize(
    ("quantitative", "qualitative"),
    [(True, True), (False, True), (True, False), (False, False)],
)
def test_programme_output_writes_combined_outcomes_sheet(
    tmp_path, dynamic, quantitative, qualitative
):
    schema = defaultdict(list)
    for row in custom_hierarchical.build_programme_hierarchical_prompt_rows():
        schema[row["class"]].append(
            {
                "attribute": row["attribute"],
                "prompt": row["prompt"],
                "datatype": row["datatype"],
            }
        )

    programme = _programme()
    if not quantitative:
        programme.quantitative_outcomes = []
    if not qualitative:
        programme.qualitative_outcomes = []
    if dynamic:
        custom_hierarchical._write_dynamic_study_outputs(
            programme,
            dict(schema),
            "Programme",
            tmp_path,
            "evaluation",
            "20260925_120000",
            "",
            export_json_file=True,
            export_xlsx_file=True,
        )
        payload = json.loads(
            (tmp_path / "evaluation_20260925_120000.json").read_text(encoding="utf-8")
        )
        assert payload == programme.model_dump()
    else:
        export_programme_csv(
            programme, "evaluation", tmp_path, "20260925_120000", write_xlsx=True
        )

    output_dir = tmp_path / "evaluation"
    assert {
        path.name for path in output_dir.glob("*.csv")
    } == {
        "programme_20260925_120000.csv",
        "interventions_20260925_120000.csv",
        "outcomes_20260925_120000.csv",
    }
    outcomes = [
        *programme.quantitative_outcomes, *programme.qualitative_outcomes
    ]
    expected_rows = [
        {field: row.model_dump().get(field, "") for field in Programme.outcome_csv_fieldnames()}
        for row in outcomes
    ]
    with (output_dir / "outcomes_20260925_120000.csv").open(
        encoding="utf-8-sig"
    ) as csv_file:
        reader = csv.DictReader(csv_file)
        assert reader.fieldnames == Programme.outcome_csv_fieldnames()
        assert list(reader) == expected_rows
    workbook = load_workbook(
        output_dir / "study_20260925_120000.xlsx", read_only=True
    )
    try:
        assert workbook.sheetnames == ["programme", "interventions", "outcomes"]
        sheet_rows = list(workbook["outcomes"].values)
        assert list(sheet_rows[0]) == Programme.outcome_csv_fieldnames()
        assert sheet_rows[1:] == [
            tuple(row[field] or None for field in Programme.outcome_csv_fieldnames())
            for row in expected_rows
        ]
    finally:
        workbook.close()


def test_export_programme_csv_writes_three_tables(tmp_path):
    export_programme_csv(
        _programme(),
        "evaluation",
        tmp_path,
        "20260925_120000",
    )

    output_dir = tmp_path / "evaluation"
    expected_headers = {
        "programme": ["countries", "responsible_author", "budget"],
        "interventions": [
            "name", "description", "time_frame", "resources_used",
            "intended_results", "source_grounding",
        ],
        "outcomes": [
            "outcome_type",
            "name",
            "description",
            "intervention_name",
            "unit_of_analysis",
            "SDG_categories",
            "results_data_source",
            "outcome_category_classification",
            "barrier",
            "barrier_source",
            "enabler",
            "enabler_source",
            "lessons_learned",
            "lessons_learned_source",
            "achievements",
            "achievements_source",
            "recommendation",
            "recommendation_source",
            "baseline_data",
            "baseline_data_source",
            "results",
        ],
    }
    assert len(list(output_dir.glob("*.csv"))) == 3
    for table_name, headers in expected_headers.items():
        output_path = output_dir / f"{table_name}_20260925_120000.csv"
        with output_path.open(encoding="utf-8-sig") as csv_file:
            reader = csv.DictReader(csv_file)
            assert reader.fieldnames == headers
            if table_name == "outcomes":
                row = next(reader)
                assert row == _programme().quantitative_outcomes[0].to_csv_row()
                assert row["results"] == (
                    "80% of partners improved delivery capacity"
                )
                assert row["outcome_type"] == "quantitative"
                row = next(reader)
                assert row["outcome_type"] == "qualitative"
                assert row["baseline_data"] == ""
                assert row["results"] == (
                    "Interviewed partners reported stronger delivery capacity"
                )


def test_new_outcome_fields_are_backwards_compatible():
    data = _programme().quantitative_outcomes[0].model_dump()
    for field in (
        "baseline_data", "baseline_data_source", "results_data_source",
        "barrier_source", "enabler_source", "lessons_learned_source",
        "achievements_source", "recommendation_source", "SDG_categories",
        "outcome_category_classification",
    ):
        del data[field]
    outcome = Quantitative_Outcome.model_validate(data)
    assert outcome.baseline_data == ""
    for field in (
        "baseline_data_source", "results_data_source", "barrier_source",
        "enabler_source", "lessons_learned_source", "achievements_source",
        "recommendation_source",
    ):
        assert getattr(outcome, field) == ""
    assert outcome.SDG_categories == ""
    assert outcome.outcome_category_classification == ""
    intervention_data = _programme().interventions[0].model_dump()
    del intervention_data["intended_results"]
    del intervention_data["source_grounding"]
    intervention = Intervention.model_validate(intervention_data)
    assert intervention.intended_results == ""
    assert intervention.source_grounding == ""


@pytest.mark.parametrize(
    "baseline",
    [
        "",
        "Before the intervention, coverage was 20% in 2023.",
        "Before the intervention, partners reported limited delivery capacity.",
        "Coverage was 20% in 2023; partners reported limited delivery capacity.",
    ],
)
def test_baseline_data_preserved_in_serialization(baseline):
    data = _programme().quantitative_outcomes[0].model_dump()
    data["baseline_data"] = baseline
    outcome = Quantitative_Outcome.model_validate(data)
    for row in (outcome.model_dump(), outcome.to_csv_row()):
        assert row["baseline_data"] == baseline
        fields = list(row)
        assert fields.index("baseline_data") + 2 == (
            fields.index("results")
        )


@pytest.mark.parametrize(
    "classification",
    ["", *SDG_CATEGORIES, "; ".join(SDG_CATEGORIES)],
)
def test_sdg_categories_accepts_exact_labels(classification):
    assert validate_sdg_categories([classification]) == [classification]
    data = _programme().quantitative_outcomes[0].model_dump()
    data["SDG_categories"] = classification
    outcome = Quantitative_Outcome.model_validate(data)
    assert outcome.to_csv_row()["SDG_categories"] == classification
    assert outcome.model_dump()["SDG_categories"] == classification


@pytest.mark.parametrize(
    "classification",
    [
        "SDG 1",
        "No Poverty",
        "No poverty; No poverty",
        "No poverty;",
        "Peace, justice and strong institutions",
        "Unknown goal",
    ],
)
def test_sdg_categories_removes_invented_or_duplicate_labels(classification):
    with pytest.warns(UserWarning, match="SDG_categories"):
        result = validate_sdg_categories([classification])
    assert result == (["No poverty"] if "No poverty;" in classification else [""])


def test_sdg_prompt_contains_exact_supplied_labels():
    for model in (Quantitative_Outcome, Qualitative_Outcome):
        description = model.model_fields["SDG_categories"].description
        assert len(SDG_CATEGORIES) == 17
        assert description.endswith("; ".join(SDG_CATEGORIES))


def test_bundled_framework_has_all_supplied_indicators_and_scopes():
    rows = list(csv.DictReader(load_indicator_framework().splitlines()))
    expected_ids = {
        f"{goal}.{indicator}"
        for goal, count in ((1, 23), (2, 8), (3, 9), (4, 9), (5, 5))
        for indicator in range(1, count + 1)
    }
    assert {row["category_id"] for row in rows} == expected_ids
    assert len(rows) == len(expected_ids)
    assert len({row["category_group"] for row in rows}) == 5
    assert all(row[column].strip() for row in rows for column in FRAMEWORK_COLUMNS)


@pytest.mark.parametrize(
    "classification",
    [
        "",
        "1.1: Percentage of pregnant women receiving at least four antenatal visits",
        "1.1: Percentage of pregnant women receiving at least four antenatal visits; 1.3: Percentage of (a) mothers and (b) newborns receiving postnatal care",
    ],
)
def test_classification_accepts_exact_entries(classification):
    assert validate_classifications(
        [classification], load_indicator_framework()
    ) == [classification]


@pytest.mark.parametrize(
    "classification",
    [
        "99.1: Invented indicator",
        "1.1: Altered statement",
        "1.1",
        "1.1: Percentage of pregnant women receiving at least four antenatal visits;",
        "1.1: Percentage of pregnant women receiving at least four antenatal visits; 1.1: Percentage of pregnant women receiving at least four antenatal visits",
    ],
)
def test_classification_removes_invalid_entries(classification):
    with pytest.warns(UserWarning, match="outcome_category_classification"):
        result = validate_classifications([classification], load_indicator_framework())
    valid = "1.1: Percentage of pregnant women receiving at least four antenatal visits"
    assert result == ([valid] if classification.startswith(valid) else [""])


def test_classification_cleanup_retains_order_and_warns_in_logs():
    first = "1.1: Percentage of pregnant women receiving at least four antenatal visits"
    second = (
        "1.3: Percentage of (a) mothers and (b) newborns receiving postnatal care"
    )
    with patch(
        "deet.hierarchical_mvp.programme_framework.logger.warning"
    ) as log_warning, pytest.warns(UserWarning):
        classifications = validate_classifications(
            iter([f"{second}; unknown; {first}; {second}", ""]),
            load_indicator_framework(),
        )
        sdgs = validate_sdg_categories(
            iter(["Climate action; Unknown; No poverty; Climate action", ""])
        )
    assert classifications == [f"{second}; {first}", ""]
    assert sdgs == ["Climate action; No poverty", ""]
    assert log_warning.call_count == 2


@pytest.mark.parametrize(
    "body",
    [
        "",
        "X,Goal,Statement,Included,\n",
        "X,Goal,Statement,Included,Excluded\nX,Goal,Other,Included,Excluded\n",
        "X,Goal,Statement; other,Included,Excluded\n",
        "X,Goal,Statement,Included,Excluded,Extra\n",
    ],
)
def test_framework_rejects_incomplete_or_ambiguous_rows(tmp_path, body):
    path = tmp_path / "invalid.csv"
    path.write_text(",".join(FRAMEWORK_COLUMNS) + "\n" + body, encoding="utf-8")
    with pytest.raises(ValueError, match="framework|category|Category"):
        load_indicator_framework(path)


def test_framework_rejects_wrong_headers_and_missing_file(tmp_path):
    path = tmp_path / "framework.csv"
    with pytest.raises(FileNotFoundError):
        load_indicator_framework(path)
    path.write_text("Category,Description\nX,Test\n", encoding="utf-8")
    with pytest.raises(ValueError, match="CSV columns"):
        load_indicator_framework(path)


@pytest.mark.parametrize("dynamic", [False, True])
@pytest.mark.parametrize("batch", [False, True])
@pytest.mark.parametrize("study_type", ["Programme", "ProgrammeWithCaching"])
def test_framework_config_reaches_single_and_batch_pipelines(
    dynamic, batch, study_type, tmp_path, monkeypatch
):
    framework_path = tmp_path / "framework.csv"
    framework_path.write_text(load_indicator_framework(), encoding="utf-8")
    report = tmp_path / "report.md"
    report.write_text("Evaluation report", encoding="utf-8")
    config = {
        "study_type": study_type,
        "llm_model": "test/model",
        "max_tokens": 1000,
        "dspy_cache": False,
        "output_parent_dir": str(tmp_path / "output"),
        "classification_framework_path": "framework.csv",
    }
    if batch:
        config["input_folder"] = str(tmp_path)
    else:
        config["input_paths"] = [str(report)]
    config_path = tmp_path / "config.json"
    config_path.write_text(json.dumps(config), encoding="utf-8")
    received_paths = []

    class FakePipeline:
        def __init__(self, classification_framework_path=None) -> None:
            received_paths.append(classification_framework_path)

        def __call__(self, **kwargs) -> Programme:
            return _programme()

    module = custom_hierarchical if dynamic else main_hierarchical
    monkeypatch.setattr(module, "configure_lm", lambda *_args, **_kwargs: None)
    if dynamic:
        schema_path = tmp_path / "schema.csv"
        schema_path.write_text("placeholder", encoding="utf-8")
        monkeypatch.setattr(module, "_load_prompt_schema", lambda _path: {})
        monkeypatch.setattr(
            module, "_build_dynamic_pipeline_for_study_type",
            lambda *_args: FakePipeline,
        )
        monkeypatch.setattr(module, "load_study_context", lambda _paths: "report")
        monkeypatch.setattr(
            module, "_write_dynamic_study_outputs",
            lambda *_args, **_kwargs: tmp_path / "output.json",
        )
        if batch:
            module.run_dynamic_batch_extraction_from_csv_schema(schema_path, config_path)
        else:
            module.run_dynamic_extraction_from_csv_schema(schema_path, config_path)
    else:
        pipeline_name = (
            "ProgrammeWithCachingExtractionPipeline"
            if study_type == "ProgrammeWithCaching"
            else "ProgrammeExtractionPipeline"
        )
        monkeypatch.setattr(module, pipeline_name, FakePipeline)
        monkeypatch.setattr(module, "read_concatenade_mds", lambda _paths: "report")
        monkeypatch.setattr(module, "save_data", lambda **_kwargs: None)
        if batch:
            module.run_predict_batch(str(config_path))
        else:
            module.run_predict_single_study(str(config_path))
    assert received_paths == [framework_path]


def test_csv_preserves_quoted_commas_quotes_multiline_text_and_bom(tmp_path):
    row = [
        "custom.10",
        "Education, health",
        'Reported "improvement", across services',
        "Measured change,\nincluding qualitative evidence",
        'Training alone, or "plans"',
    ]
    content = StringIO(newline="")
    writer = csv.writer(content)
    writer.writerow(FRAMEWORK_COLUMNS)
    writer.writerow(row)
    path = tmp_path / "custom.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        handle.write(content.getvalue())

    framework = load_indicator_framework(path)

    assert list(csv.reader(StringIO(framework)))[1] == row
    validate_classifications([f"{row[0]}: {row[2]}"], framework)


@pytest.mark.parametrize("value", [None, "", "  ", 42])
def test_framework_config_rejects_invalid_paths(tmp_path, value):
    with pytest.raises(TypeError, match="non-empty string"):
        resolve_framework_config(
            {"study_type": "Programme", "classification_framework_path": value},
            tmp_path / "config.json",
        )


def test_framework_config_rejects_other_study_types(tmp_path):
    with pytest.raises(ValueError, match="only"):
        resolve_framework_config(
            {"study_type": "RCT", "classification_framework_path": "framework.csv"},
            tmp_path / "config.json",
        )


def test_framework_config_without_override_uses_default(tmp_path):
    assert resolve_framework_config(
        {"study_type": "Programme"}, tmp_path / "config.json"
    ) is None

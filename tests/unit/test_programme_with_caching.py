import csv
import json
from collections import defaultdict
from os.path import commonprefix
from unittest.mock import patch

import dspy
import pytest
from dspy.adapters.json_adapter import _get_structured_outputs_response_format
from dspy.utils import DummyLM
from litellm import ModelResponse
from openpyxl import load_workbook

from deet import custom_hierarchical, main_hierarchical
from deet.hierarchical_mvp.extraction_logging import extraction_usage_log
from deet.hierarchical_mvp.ProgrammeModel import (
    Programme,
    Qualitative_Outcome,
    Quantitative_Outcome,
)
from deet.hierarchical_mvp.ProgrammeWithCachingExtraction import (
    ProgrammeWithCachingExtractionPipeline,
)
from deet.hierarchical_mvp.ProgrammeWithCachingModel import ProgrammeWithCaching

STAGES = ("programme", "quantitative", "qualitative")


@pytest.fixture
def programme():
    quantitative = dict.fromkeys(Quantitative_Outcome.model_fields, "")
    quantitative.update(
        outcome_type="quantitative",
        name="Improved skills",
        intervention_name="Training",
        results="80% improved",
        baseline_data="NR",
        SDG_categories="Quality education",
    )
    qualitative = dict.fromkeys(Qualitative_Outcome.model_fields, "")
    qualitative.update(
        outcome_type="qualitative",
        name="Confidence",
        intervention_name="Training",
        results="Interviewees reported increased confidence",
    )
    return ProgrammeWithCaching.model_validate(
        {
            "programme_characteristics": {
                "countries": "Kenya",
                "responsible_author": "UN",
                "budget": "",
            },
            "interventions": [
                {
                    "name": "Training",
                    "description": "Skills training",
                    "time_frame": "2025",
                    "resources_used": "",
                }
            ],
            "quantitative_outcomes": [quantitative],
            "qualitative_outcomes": [qualitative],
        }
    )


@pytest.fixture(params=[False, True], ids=["predefined", "custom"])
def pipeline(request):
    if not request.param:
        return ProgrammeWithCachingExtractionPipeline()
    schema = defaultdict(list)
    for row in custom_hierarchical.build_programme_hierarchical_prompt_rows():
        schema[row["class"]].append(row)
    schema["Qualitative_Outcome"].append(
        {
            "attribute": "participant_group",
            "prompt": "The participants describing their experiences.",
            "datatype": "str",
        }
    )
    pipeline_class = custom_hierarchical._build_dynamic_pipeline_for_study_type(  # noqa: SLF001
        "ProgrammeWithCaching", dict(schema)
    )
    return pipeline_class()


def _stage_answers(pipeline, programme):
    programme = pipeline.programme_model.model_validate(programme.model_dump())
    return [
        pipeline.step_model(
            programme_characteristics=programme.programme_characteristics,
            interventions=programme.interventions,
            quantitative_outcomes=[],
            qualitative_outcomes=[],
        ),
        pipeline.step_model(
            programme_characteristics=None,
            interventions=[],
            quantitative_outcomes=programme.quantitative_outcomes,
            qualitative_outcomes=[],
        ),
        pipeline.step_model(
            programme_characteristics=None,
            interventions=[],
            quantitative_outcomes=[],
            qualitative_outcomes=programme.qualitative_outcomes,
        ),
    ]


@pytest.mark.parametrize("adapter_type", [dspy.ChatAdapter, dspy.JSONAdapter])
@pytest.mark.parametrize(
    ("quantitative", "qualitative"),
    [
        (True, True),
        (True, False),
        (False, True),
        (False, False),
    ],
)
def test_pipeline_parses_outputs_and_shares_full_report_prefix(
    pipeline, programme, adapter_type, quantitative, qualitative
):
    if not quantitative:
        programme.quantitative_outcomes = []
    if not qualitative:
        programme.qualitative_outcomes = []
    answers = _stage_answers(pipeline, programme)
    adapter = adapter_type()
    lm = DummyLM(
        [{"extraction": answer.model_dump(mode="json")} for answer in answers],
        adapter=adapter,
    )
    report = "BEGIN_REPORT\n" + "Detailed evaluation evidence.\n" * 5000 + "END_REPORT"
    with dspy.context(lm=lm, adapter=adapter):
        result = pipeline(context=report)
    assert (
        result.model_dump()
        == pipeline.programme_model.model_validate(programme.model_dump()).model_dump()
    )
    assert len(lm.history) == len(pipeline.predictors()) == 3
    messages = [call["messages"] for call in lm.history]
    assert messages[0][0] == messages[1][0] == messages[2][0]
    shared = commonprefix([call[-1]["content"] for call in messages])
    assert shared.count(report) == 1
    assert pipeline.classification_framework in shared
    assert shared.index(report) < shared.index(pipeline.classification_framework)
    assert len(shared) > len(report) + len(pipeline.classification_framework)
    for stage, call in zip(STAGES, messages, strict=True):
        content = call[-1]["content"]
        assert (
            content.index("END_REPORT")
            < content.index("[[ ## extraction_stage ## ]]")
            < content.index("[[ ## previous_extraction ## ]]")
        )
        assert f"[[ ## extraction_stage ## ]]\n{stage}" in content
    last_request = messages[2][-1]["content"]
    assert "Training" in last_request
    if quantitative:
        assert "80% improved" in last_request
    else:
        assert "80% improved" not in last_request
    if pipeline.programme_model is ProgrammeWithCaching:
        assert isinstance(result, Programme)
    else:
        assert (
            "participant_group"
            in (
                pipeline.step_model.model_json_schema()["$defs"][
                    "DynamicQualitative_Outcome"
                ]["properties"]
            )
        )


def test_native_response_schema_is_identical_for_all_passes(pipeline):
    schemas = [
        _get_structured_outputs_response_format(predictor.signature).model_json_schema()
        for predictor in pipeline.predictors()
    ]
    assert schemas[0] == schemas[1] == schemas[2]
    assert set(schemas[0]["$defs"][pipeline.step_model.__name__]["properties"]) == {
        "programme_characteristics",
        "interventions",
        "quantitative_outcomes",
        "qualitative_outcomes",
    }


def test_openai_request_shapes_and_stage_logs_with_custom_framework(
    pipeline, programme, tmp_path
):
    framework = (
        "category_id,category_group,category_statement,inclusion_criteria,"
        "exclusion_criteria\n"
        "X.1,Capacity,Delivery capacity,Improved skills,Plans only\n"
    )
    framework_path = tmp_path / "framework.csv"
    framework_path.write_text(framework, encoding="utf-8")
    pipeline = type(pipeline)(classification_framework_path=framework_path)
    programme.quantitative_outcomes[
        0
    ].outcome_category_classification = "X.1: Delivery capacity"
    answers = _stage_answers(pipeline, programme)
    responses = [
        ModelResponse(
            model="gpt-4.1",
            choices=[
                {
                    "message": {
                        "role": "assistant",
                        "content": json.dumps(
                            {"extraction": answer.model_dump(mode="json")}
                        ),
                    }
                }
            ],
            usage={"prompt_tokens": 2000, "completion_tokens": 100},
        )
        for answer in answers
    ]
    lm = dspy.LM("openai/gpt-4.1", cache=False)
    with (
        dspy.context(lm=lm, adapter=dspy.JSONAdapter()),
        patch.object(lm, "forward", side_effect=responses) as forward,
        extraction_usage_log(tmp_path, "evaluation", "ProgrammeWithCaching", lm.model),
    ):
        result = pipeline(context="ORIGINAL_REPORT")
    assert result.quantitative_outcomes[0].outcome_category_classification == (
        "X.1: Delivery capacity"
    )
    requests = [call.kwargs for call in forward.call_args_list]
    assert len(requests) == 3
    schemas = [request["response_format"].model_json_schema() for request in requests]
    assert schemas[0] == schemas[1] == schemas[2]
    assert (
        requests[0]["messages"][0]
        == requests[1]["messages"][0]
        == (requests[2]["messages"][0])
    )
    shared = commonprefix([request["messages"][-1]["content"] for request in requests])
    assert "ORIGINAL_REPORT" in shared
    assert framework in shared
    log = json.loads(next(tmp_path.glob("extraction_log_*.json")).read_text())
    assert [call["stage"] for call in log["calls"]] == [
        "ExtractProgrammeWithCachingInfo",
        "ExtractProgrammeWithCachingQuantitative",
        "ExtractProgrammeWithCachingQualitative",
    ]
    assert [call["attempt"] for call in log["calls"]] == [1, 1, 1]


def test_pipeline_cleans_outcomes_before_passing_to_next_stage(pipeline, programme):
    programme.quantitative_outcomes[0].SDG_categories = "Unknown; Quality education"
    programme.qualitative_outcomes[0].outcome_category_classification = "Invalid"
    answers = _stage_answers(pipeline, programme)
    lm = DummyLM([{"extraction": answer.model_dump(mode="json")} for answer in answers])
    with (
        dspy.context(lm=lm),
        pytest.warns(UserWarning, match="removed invalid or duplicate") as warnings,
    ):
        result = pipeline(context="report")
    assert len(warnings) == 2
    assert result.quantitative_outcomes[0].SDG_categories == "Quality education"
    assert result.qualitative_outcomes[0].outcome_category_classification == ""
    assert "Unknown" not in lm.history[2]["messages"][-1]["content"]


def test_missing_programme_metadata_fails_explicitly(pipeline):
    empty = pipeline.step_model(
        programme_characteristics=None,
        interventions=[],
        quantitative_outcomes=[],
        qualitative_outcomes=[],
    )
    lm = DummyLM([{"extraction": empty.model_dump()}])
    with (
        dspy.context(lm=lm),
        pytest.raises(ValueError, match="must return programme_characteristics"),
    ):
        pipeline(context="report")


def test_unrequested_stage_results_fail_explicitly(pipeline, programme):
    answers = _stage_answers(pipeline, programme)
    answers[0].quantitative_outcomes = answers[1].quantitative_outcomes
    lm = DummyLM([{"extraction": answers[0].model_dump(mode="json")}])
    with (
        dspy.context(lm=lm),
        pytest.raises(
            ValueError,
            match="programme stage returned unexpected quantitative_outcomes",
        ),
    ):
        pipeline(context="report")


def test_main_dispatches_separate_caching_pipeline(programme, tmp_path):
    framework_path = tmp_path / "framework.csv"
    with patch.object(
        main_hierarchical, "ProgrammeWithCachingExtractionPipeline"
    ) as pipeline_class:
        pipeline_class.return_value.return_value = programme
        result = main_hierarchical.extract(
            "report", "ProgrammeWithCaching", framework_path
        )
    pipeline_class.assert_called_once_with(classification_framework_path=framework_path)
    pipeline_class.return_value.assert_called_once_with(context="report")
    assert result is programme


def test_custom_prompt_csv_uses_existing_programme_fields(tmp_path):
    path = custom_hierarchical.write_hierarchical_prompts_csv(
        study_type="ProgrammeWithCaching", csv_outpath=tmp_path / "prompts.csv"
    )
    with path.open(encoding="utf-8-sig") as file:
        rows = list(csv.DictReader(file))
    assert rows == custom_hierarchical.build_programme_hierarchical_prompt_rows()


@pytest.mark.parametrize("dynamic", [False, True])
def test_caching_pipeline_exports_compatible_files(programme, tmp_path, dynamic):
    if dynamic:
        schema = defaultdict(list)
        for row in custom_hierarchical.build_programme_hierarchical_prompt_rows():
            schema[row["class"]].append(row)
        custom_hierarchical._write_dynamic_study_outputs(  # noqa: SLF001
            programme,
            dict(schema),
            "ProgrammeWithCaching",
            tmp_path,
            "evaluation",
            "test",
            "",
            export_json_file=True,
            export_xlsx_file=True,
        )
    else:
        main_hierarchical.save_data(
            programme,
            ["evaluation.md"],
            str(tmp_path),
            study_type="ProgrammeWithCaching",
            export_json_file=True,
            export_xlsx_file=True,
        )
    payload = json.loads(next(tmp_path.glob("*.json")).read_text())
    assert payload == programme.model_dump()
    assert len(list((tmp_path / "evaluation").glob("*.csv"))) == 3
    workbook = load_workbook(
        next((tmp_path / "evaluation").glob("*.xlsx")), read_only=True
    )
    try:
        assert workbook.sheetnames == ["programme", "interventions", "outcomes"]
    finally:
        workbook.close()
